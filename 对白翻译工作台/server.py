# -*- coding: utf-8 -*-
"""
对白翻译工作台 —— 本地小后端
  GET  /                  index.html
  GET  /api/data          当前剧本：每句(rec/alt/选择)
  GET  /api/scripts       列出可导入 docx 及翻译状态
  POST /api/choice        {idx, text?} 或 {idx, skip:true/false}
  POST /api/load          {path} 切换到某剧本
  POST /api/translate     {path} 未翻译则先跑双候选翻译，再载入
  POST /api/export        导出 中上英下.docx + 纯英.txt

用法: python server.py [--docx 剧本.docx] [--port 8100]
choices/缓存/成品 都存在剧本所在目录。
"""
import os, sys, re, json, copy, shutil, subprocess, argparse
sys.stdout.reconfigure(encoding='utf-8')
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT     = os.path.dirname(BASE_DIR)                       # E:\desk
DEFAULT_DOCX = os.path.join(ROOT, '鲨鱼剧本AI测试.docx')
TRANSLATE_PY = os.path.join(ROOT, 'translate_docx.py')
SCAN_DIRS = [ROOT,
             os.path.join(ROOT, '剧本(claude如果查剧本点这里）'),
             os.path.join(ROOT, 'AI work'),
             os.path.join(ROOT, 'AI download')]

def to_num(t):
    t = t.strip()
    if t.isdigit(): return int(t)
    cn = {'一':1,'二':2,'三':3,'四':4,'五':5,'六':6,'七':7,'八':8,'九':9,'十':10}
    if t in cn: return cn[t]
    if '十' in t:
        a, _, b = t.partition('十')
        return (cn.get(a, 1) if a else 1) * 10 + (cn.get(b, 0) if b else 0)
    return None

def is_dialogue(s):
    s = s.strip()
    if not s or '：' not in s: return False
    head = s.split('：', 1)[0]
    if not head or head in ('人物','时间','地点','场景','字幕','标题'): return False
    if s[0] in '△【(（［#': return False
    if head.startswith('人物') or len(head) > 24: return False
    return bool(s.split('：', 1)[1].strip())

def insert_english_after(p_el, text):
    new_el = copy.deepcopy(p_el)
    for child in list(new_el):
        if child.tag != qn('w:pPr'):
            new_el.remove(child)
    rpr = None
    runs = p_el.findall(qn('w:r'))
    if runs: rpr = runs[0].find(qn('w:rPr'))
    r = OxmlElement('w:r')
    if rpr is not None: r.append(copy.deepcopy(rpr))
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); new_el.append(r); p_el.addnext(new_el)
    return new_el

def cache_state(base):
    if os.path.isfile(base + '_candidates_cache.json'): return 'candidates'
    if os.path.isfile(base + '_cache.json'):            return 'single'
    return 'none'

def count_episodes(path):
    try:
        doc = Document(path)
        return sum(1 for p in doc.paragraphs
                   if re.search(r'第\s*[0-9一二三四五六七八九十百千〇零]+\s*[集回]', p.text))
    except Exception:
        return None

class Store:
    def __init__(self, docx):
        self.docx = os.path.abspath(docx)
        base = os.path.splitext(self.docx)[0]
        self.choices_path = base + '_choices.json'
        doc = Document(self.docx)
        self.paras = [x.text for x in doc.paragraphs]
        cur, self.ep = 1, []
        for i, p in enumerate(self.paras):
            m = re.search(r'第\s*([0-9一二三四五六七八九十百千〇零]+)\s*[集回]', p)
            if m:
                nn = to_num(m.group(1))
                if nn: cur = nn
            self.ep.append(cur)
        self.dial = [{'idx': i, 'ep': self.ep[i], 'zh': self.paras[i].split('：', 1)[1].strip()}
                     for i, p in enumerate(self.paras) if is_dialogue(p)]
        by_ep = {}
        for dl in self.dial: by_ep.setdefault(dl['ep'], []).append(dl['idx'])
        self.rec, self.alt = {}, {}
        for cand in (base + '_candidates_cache.json', base + '_cache.json'):
            if not os.path.isfile(cand): continue
            c = json.load(open(cand, encoding='utf-8'))
            two_ep = {}
            for e, idxs in by_ep.items():
                arr = c.get(str(e))
                if isinstance(arr, list) and len(arr) == 2 * len(idxs):
                    two_ep[e] = True
                elif isinstance(arr, list) and len(arr) == len(idxs):
                    two_ep[e] = False
            if not two_ep: continue
            for e, idxs in by_ep.items():
                if e not in two_ep: continue
                arr = c[str(e)]
                two = two_ep[e]
                for k, idx in enumerate(idxs):
                    self.rec[idx] = arr[2 * k] if two else arr[k]
                    if two: self.alt[idx] = arr[2 * k + 1]
            break
        self.choices = {}
        if os.path.isfile(self.choices_path):
            self.choices = json.load(open(self.choices_path, encoding='utf-8'))

    def save_choices(self):
        with open(self.choices_path, 'w', encoding='utf-8') as f:
            json.dump(self.choices, f, ensure_ascii=False, indent=1)

    def state(self, dl):
        idx = str(dl['idx'])
        rec = self.rec.get(dl['idx'], '')
        alt = self.alt.get(dl['idx'], '')
        c = self.choices.get(idx)
        if isinstance(c, dict) and c.get('skip'):
            return {'skip': True, 'rec': rec, 'alt': alt, 'text': '', 'src': 'skip'}
        if isinstance(c, str) and c:
            return {'skip': False, 'rec': rec, 'alt': alt, 'text': c,
                    'src': 'A' if c == rec else ('B' if c == alt else 'custom')}
        return {'skip': False, 'rec': rec, 'alt': alt, 'text': rec, 'src': 'A' if rec else 'none'}

    def data(self):
        eps = {}
        for dl in self.dial:
            eps.setdefault(dl['ep'], []).append({'idx': dl['idx'], 'zh': dl['zh'], **self.state(dl)})
        return {'docx': os.path.basename(self.docx),
                'path': self.docx,
                'episodes': [{'ep': e, 'lines': eps[e]} for e in sorted(eps)]}

    def export(self):
        chosen = {}
        for dl in self.dial:
            st = self.state(dl)
            if not st.get('skip') and st.get('text'):
                chosen[dl['idx']] = st['text'].strip()
        base = os.path.splitext(self.docx)[0]
        out_docx = base + '_成品_中上英下.docx'
        shutil.copyfile(self.docx, out_docx)
        d2 = Document(out_docx)
        for i, p in enumerate(d2.paragraphs):
            if i in chosen:
                insert_english_after(p._p, chosen[i])
        d2.save(out_docx)
        out_txt = base + '_成品_纯英.txt'
        lines, last_ep = [], None
        for dl in self.dial:
            if dl['idx'] not in chosen: continue
            if dl['ep'] != last_ep:
                if last_ep is not None: lines.append('')
                lines.append(f'[第{dl["ep"]}集]')
                last_ep = dl['ep']
            lines.append(chosen[dl['idx']])
        with open(out_txt, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))
        return {'docx': out_docx, 'txt': out_txt, 'lines': len(chosen)}

STORE = None

def load_store(path):
    global STORE
    STORE = Store(path)

def list_scripts():
    out, seen = [], set()
    for d in SCAN_DIRS:
        if not os.path.isdir(d): continue
        for f in sorted(os.listdir(d)):
            if not f.lower().endswith('.docx'): continue
            if any(k in f for k in ('成品', '中上英下', '双候选', '~$')): continue
            p = os.path.join(d, f)
            if p in seen: continue
            seen.add(p)
            base = os.path.splitext(p)[0]
            out.append({'name': f, 'path': p,
                        'ep': count_episodes(p),
                        'state': cache_state(base)})
    return out

class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def _send(self, code, obj, ctype='application/json; charset=utf-8'):
        body = obj if isinstance(obj, bytes) else json.dumps(obj, ensure_ascii=False).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', ctype)
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)
    def do_GET(self):
        p = urlparse(self.path).path
        if p in ('/', '/index.html'):
            raw = open(os.path.join(BASE_DIR, 'index.html'), encoding='utf-8').read()
            self._send(200, raw.encode('utf-8'), 'text/html; charset=utf-8')
        elif p == '/api/data':
            self._send(200, STORE.data())
        elif p == '/api/scripts':
            self._send(200, {'scripts': list_scripts()})
        else:
            self._send(404, {'error': 'not found'})
    def do_POST(self):
        try:
            self._post()
        except Exception as e:
            try: self._send(500, {'error': str(e)})
            except Exception: pass
    def _post(self):
        p = urlparse(self.path).path
        ln = int(self.headers.get('content-length') or 0)
        try:
            body = json.loads(self.rfile.read(ln).decode('utf-8', 'replace') or '{}')
        except Exception:
            body = {}
        if p == '/api/choice':
            idx = str(body['idx'])
            if body.get('skip'):
                STORE.choices[idx] = {'skip': True}
            else:
                STORE.choices[idx] = body.get('text', '')
            STORE.save_choices()
            self._send(200, {'ok': True})
        elif p == '/api/load':
            try:
                load_store(body['path'])
                self._send(200, {'ok': True, 'data': STORE.data()})
            except Exception as e:
                self._send(500, {'error': str(e)})
        elif p == '/api/translate':
            path = body['path']
            base = os.path.splitext(path)[0]
            try:
                if cache_state(base) != 'candidates':
                    r = subprocess.run([sys.executable, TRANSLATE_PY, '--docx', path, '--candidates'],
                                       capture_output=True, text=True, timeout=3600)
                    if r.returncode != 0:
                        return self._send(500, {'error': r.stderr[-800:] or 'translate failed'})
                load_store(path)
                self._send(200, {'ok': True, 'data': STORE.data()})
            except Exception as e:
                self._send(500, {'error': str(e)})
        elif p == '/api/export':
            try:
                self._send(200, STORE.export())
            except Exception as e:
                self._send(500, {'error': str(e)})
        else:
            self._send(404, {'error': 'not found'})

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--docx', default=DEFAULT_DOCX)
    ap.add_argument('--port', type=int, default=8100)
    a = ap.parse_args()
    load_store(a.docx)
    srv = ThreadingHTTPServer(('127.0.0.1', a.port), H)
    print(f'对白翻译工作台: http://127.0.0.1:{a.port}')
    print(f'剧本: {STORE.docx} | 台词 {len(STORE.dial)} 句')
    srv.serve_forever()
